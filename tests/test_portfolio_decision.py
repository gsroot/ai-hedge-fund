from __future__ import annotations

import copy
import io
import json
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / ".agents/skills/portfolio-report/scripts"
sys.path.insert(0, str(SCRIPTS))
from portfolio_decision import digest, load_snapshot, save_snapshot, validate_decision
from decision_report import render_text, report_rows, write_reports
from evaluate_decision import evaluate, save_evaluation


def fixture():
    return json.loads((ROOT / "tests/fixtures/portfolio_decision.json").read_text(encoding="utf-8"))


def validate(f):
    return validate_decision(f["decision"], f["predict"], f["analyses"]["analyses"], f["risk"], ["buffett", "lynch", "fisher"])


class DecisionTests(unittest.TestCase):
    def test_actions_use_known_holdings_without_order_quantities(self):
        f = fixture()
        result = validate(f)
        self.assertEqual(result["actions"]["AAA"], {"action": "target_only", "target_weight": .15})
        f["decision"]["user_context"]["holdings"] = {"AAA": .2, "MARKET": .1, "OLD": .1, "CASH": .6}
        with self.assertRaisesRegex(ValueError, "explicit review"):
            validate(f)
        old = copy.deepcopy(f["decision"]["assets"][0])
        old.update(ticker="OLD", weight=0, selection_reason="Review found alternatives more attractive")
        f["decision"]["assets"].append(old)
        result = validate(f)
        self.assertEqual(result["actions"]["AAA"]["action"], "reduce")
        self.assertEqual(result["actions"]["MARKET"]["action"], "add")
        self.assertEqual(result["actions"]["OLD"]["action"], "exit")
        self.assertAlmostEqual(result["actions"]["AAA"]["change_weight"], -.05)
        self.assertNotIn("quantity", result["actions"]["AAA"])

    def test_combined_cost_cannot_make_negative_terminal_holdings(self):
        f = fixture()
        f["decision"]["base_currency"] = "KRW"
        f["decision"]["evaluation"].update(exit_cost_bps=9000, fx_cost_bps=2000)
        with self.assertRaisesRegex(ValueError, "combined exit costs"):
            validate(f)

    def test_long_source_details_are_not_truncated_by_excel(self):
        f = fixture()
        long_reason = "가상 자료 " * 8000
        f["decision"]["assets"][0]["thesis"]["rationale"] = long_reason
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.json"
            write_reports(validate(f), path, path.with_suffix(".xlsx"))
            ws = load_workbook(path.with_suffix(".xlsx")).active
            parts = [row[2] for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == "AAA" and str(row[1]).startswith("rationale [")]
            self.assertEqual("".join(parts), long_reason)

    def test_prior_only_and_partial_opinion_do_not_block_sizing(self):
        result = validate(fixture())
        self.assertEqual(result["weights"], {"AAA": .15, "MARKET": .6})
        self.assertEqual(result["cash_weight"], .25)
        self.assertEqual(result["status"], "conditional")
        self.assertEqual(result["performance_evidence"]["status"], "prior_only")
        self.assertEqual(result["evidence"]["AAA"]["independent_analyses"]["fisher"]["data_quality"], "insufficient")
        self.assertAlmostEqual(result["constraint_checks"]["total_issuers"]["AAA"], .21)
        self.assertEqual(result["decision"]["user_context"]["required_amount"], None)

    def test_missing_opinion_not_fabricated_and_dates_aligned(self):
        f = fixture()
        del f["analyses"]["analyses"]["MARKET"]["lynch"]
        with self.assertRaisesRegex(ValueError, "independent analysis missing"):
            validate(f)
        f = fixture()
        del f["analyses"]["analyses"]["AAA"]["buffett"]["analysis_source"]
        with self.assertRaisesRegex(ValueError, "analysis source must be independent"):
            validate(f)
        f = fixture()
        f["risk"]["analysis_date"] = "2026-09-07"
        with self.assertRaisesRegex(ValueError, "dates must agree"):
            validate(f)

    def test_invalid_numbers_duplicates_and_future_sources_rejected(self):
        mutations = [
            lambda d: d["assets"][0].update(weight=float("nan")),
            lambda d: d["assets"][0].update(weight=-.1),
            lambda d: d["assets"][0]["quote"].update(value=0),
            lambda d: d["assets"].append(copy.deepcopy(d["assets"][0])),
            lambda d: d["sources"][0].update(available_at="2026-09-09T20:00:00Z"),
            lambda d: d["assets"][0]["quote"].update(source_id="missing"),
            lambda d: d["sources"][0].update(retrieved_at="2026-09-10T00:00:00Z"),
        ]
        for mutate in mutations:
            f = fixture()
            mutate(f["decision"])
            with self.subTest(mutate=mutate), self.assertRaises(ValueError):
                validate(f)

    def test_no_hidden_rescaling_and_cash_rule_is_advisory(self):
        f = fixture()
        f["decision"]["assets"][1]["weight"] = .8
        f["decision"]["cash"].update(weight=.05, liquidity_weight=.05)
        f["decision"]["user_context"]["liquidity_min_weight"] = .05
        result = validate(f)
        self.assertEqual(result["weights"]["MARKET"], .8)
        self.assertEqual(result["cash_weight"], .05)
        self.assertAlmostEqual(result["market_regime"]["cash_difference"], -.05)
        f["decision"]["cash"]["weight"] = .04
        failed = validate(f)
        self.assertEqual(failed["status"], "infeasible")
        self.assertEqual(failed["weights"], {})
        self.assertIsNone(failed["cash_weight"])

    def test_liquidity_constraints_and_user_assumptions(self):
        f = fixture()
        f["decision"]["user_context"]["liquidity_basis"] = "user_confirmed"
        with self.assertRaisesRegex(ValueError, "unknown spending"):
            validate(f)
        f = fixture()
        f["decision"]["user_context"].update(required_amount=3000, portfolio_value=10000, spending_fx_to_base=1, spending_fx_basis="same currency")
        with self.assertRaisesRegex(ValueError, "contradicts required amount"):
            validate(f)
        f = fixture()
        f["decision"]["cash"]["liquidity_weight"] = .1
        self.assertEqual(validate(f)["status"], "infeasible")

    def test_typed_caps_overlap_and_incomplete_holdings(self):
        f = fixture()
        f["decision"]["constraints"]["max_total_issuer_weight"] = .2
        result = validate(f)
        self.assertEqual(result["status"], "infeasible")
        self.assertTrue(any("AAA" in s for s in result["constraint_checks"]["violations"]))
        f["decision"]["assets"][1]["etf"]["holdings"] = []
        f["decision"]["assets"][1]["etf"]["holdings_complete"] = False
        self.assertEqual(validate(f)["status"], "blocked_input")
        del f["decision"]["constraints"]["max_total_issuer_weight"]
        self.assertEqual(validate(f)["status"], "conditional")

    def test_concentrated_etf_requires_explicit_cap(self):
        f = fixture()
        f["decision"]["assets"][1]["etf"]["leverage"] = 2
        with self.assertRaisesRegex(ValueError, "explicit asset cap"):
            validate(f)
        f["decision"]["constraints"]["asset_caps"] = {"MARKET": .7}
        self.assertEqual(validate(f)["weights"]["MARKET"], .6)

    def test_etf_only_and_cash_only_are_deliberate_valid_choices(self):
        f = fixture()
        f["decision"]["assets"][0]["weight"] = 0
        f["decision"]["assets"][1]["weight"] = 1
        f["decision"]["cash"].update(weight=0, liquidity_weight=0)
        f["decision"]["user_context"]["liquidity_min_weight"] = 0
        self.assertEqual(validate(f)["weights"], {"MARKET": 1})
        f["decision"]["assets"][1]["weight"] = 0
        f["decision"]["cash"]["weight"] = 1
        self.assertEqual(validate(f)["cash_weight"], 1)

    def test_failed_excluded_candidate_does_not_become_negative_signal(self):
        for state in ("timeout", "credentials_missing", "quota", "stale", "schema"):
            f = fixture()
            f["decision"]["sources"].append({**f["decision"]["sources"][0], "id": "failure", "status": state})
            bad = {**copy.deepcopy(f["decision"]["assets"][0]), "ticker": "BAD", "weight": 0, "selection_reason": "Cannot assess price; alternatives available"}
            bad["quote"].update(source_id="failure", value=None)
            f["decision"]["assets"].append(bad)
            with self.subTest(state=state):
                result = validate(f)
                self.assertEqual(result["weights"], {"AAA": .15, "MARKET": .6})
                self.assertIn(state, result["evidence"]["BAD"]["price_issues"][0])

    def test_all_prices_failed_is_not_successful_cash_recommendation(self):
        f = fixture()
        for asset in f["decision"]["assets"]:
            asset["weight"] = 0
            asset["quote"]["value"] = None
        f["decision"]["cash"]["weight"] = 1
        f["decision"]["sources"][0]["status"] = "timeout"
        result = validate(f)
        self.assertEqual(result["status"], "blocked_input")
        self.assertIsNone(result["cash_weight"])

    def test_stale_price_and_revised_fundamentals(self):
        f = fixture()
        f["decision"]["sources"][0]["as_of"] = "2026-08-01T20:00:00Z"
        f["decision"]["assets"][0]["quote"]["as_of"] = "2026-08-01T20:00:00Z"
        self.assertEqual(validate(f)["status"], "blocked_input")
        f = fixture()
        f["decision"]["sources"][1].update(status="revision", revision_reason="Current filing corrected; not a historical vintage")
        self.assertEqual(validate(f)["status"], "conditional")

    def test_score_does_not_generate_forecast_or_change_judgment(self):
        f = fixture()
        before = validate(f)
        f["predict"]["rankings"][0]["total_score"] = 1
        self.assertEqual(validate(f)["weights"], before["weights"])
        f["decision"]["assets"][0]["thesis"]["scenarios"] = [dict(assumptions="test",calculation="test",horizon="6m",currency="USD",return_components="price only",basis="score",target_price=150,source_refs=["prices"])]
        with self.assertRaisesRegex(ValueError, "scenario basis"):
            validate(f)

    def test_round_trip_report_consistency_and_formula_safety(self):
        f = fixture()
        f["decision"]["assets"][0]["selection_reason"] = '=HYPERLINK("https://example.invalid")'
        result = validate(f)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "result.json"
            text = write_reports(result, path, path.with_suffix(".xlsx"))
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), result)
            workbook = load_workbook(path.with_suffix(".xlsx"))
            ws = workbook.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertIn(("최종 비중", "AAA", .15), rows)
            self.assertIn(("최종 비중", "CASH", .25), rows)
            self.assertIn("최종 비중 | MARKET: 60.00%", text)
            self.assertIn("현금 판단 | annual_rate: 4.00%", text)
            self.assertIn("시장 참고값 | target_cash_weight: 10.00%", text)
            self.assertNotIn("0.169999999999", text)
            self.assertIn(("판단", "decision_id", result["decision_id"]), rows)
            cash_rate_row = next(i for i, row in enumerate(rows, start=2) if row[:2] == ("현금 판단", "annual_rate"))
            self.assertEqual(ws.cell(cash_rate_row, 3).number_format, "0.00%")
            self.assertIn("1-3 years", text)
            self.assertIn("6 months", text)
            formula_cells = [c for row in ws for c in row if c.value == f["decision"]["assets"][0]["selection_reason"]]
            self.assertTrue(formula_cells)
            self.assertEqual(formula_cells[0].data_type, "s")

    def test_snapshot_id_integrity_previous_link_and_credentials(self):
        result = validate(fixture())
        with tempfile.TemporaryDirectory() as tmp:
            folder = save_snapshot(result, tmp)
            self.assertEqual(load_snapshot(folder), result)
            with self.assertRaises(FileExistsError):
                save_snapshot(result, tmp)
            followup = copy.deepcopy(result)
            followup["decision_id"] = followup["decision"]["decision_id"] = "next-decision"
            followup["decision"]["previous_decision_id"] = result["decision_id"]
            self.assertEqual(load_snapshot(save_snapshot(followup, tmp))["decision"]["previous_decision_id"], result["decision_id"])
            (folder / "decision.json").write_text("{}")
            with self.assertRaisesRegex(ValueError, "integrity"):
                load_snapshot(folder)
        for secret in ({"api_key": "test-secret"}, {"reference": "https://example.invalid?api_key=test-secret"}):
            f = fixture()
            f["decision"]["metadata"] = secret
            with self.assertRaisesRegex(ValueError, "credential"):
                validate(f)


class EvaluationTests(unittest.TestCase):
    def test_zero_cash_does_not_require_irrelevant_fx_series(self):
        f = fixture()
        f["decision"]["cash"].update(weight=0, liquidity_weight=0, currency="KRW")
        f["decision"]["assets"][1]["weight"] = .85
        f["decision"]["user_context"]["liquidity_min_weight"] = 0
        actual = evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")
        self.assertEqual(actual["status"], "complete")
        self.assertEqual(actual["contributions"]["CASH"], 0)

    def test_buy_hold_hand_calculation_costs_cash_and_negative_relative(self):
        f = fixture()
        result = validate(f)
        actual = evaluate(result, f["prices"], "2027-03-10T00:00:00Z")
        from portfolio_decision import timestamp
        days = (timestamp("2027-03-09T21:00:00Z") - timestamp("2026-09-09T13:30:00Z")).total_seconds()/86400
        cash_return = .25 * ((1.04)**(days/365.25)-1)
        expected = .15*(1.2/1.001*.999-1)+.6*(1.3/1.001*.999-1)+cash_return
        self.assertEqual(actual["status"], "complete")
        self.assertAlmostEqual(actual["portfolio_return"], expected)
        self.assertAlmostEqual(sum(actual["contributions"].values()), expected)
        self.assertLess(actual["relative_return"], 0)
        self.assertEqual(actual["thesis_reviews"]["AAA"]["status"], "not_reviewed")

    def test_no_preformation_return_or_daily_rebalance(self):
        f = fixture()
        f["decision"]["cash"]["annual_rate"] = 0
        f["decision"]["evaluation"].update(entry_cost_bps=0, exit_cost_bps=0)
        f["prices"]["series"]["AAA"]["bars"][0]["close"] = 200
        f["prices"]["series"]["AAA"]["bars"][1]["close"] = 100
        f["prices"]["series"]["MARKET"]["bars"][0]["close"] = 100
        f["prices"]["series"]["MARKET"]["bars"][1]["close"] = 100
        f["prices"]["series"]["AAA"]["bars"].insert(0, dict(date="2026-09-08", open=1, close=100))
        actual = evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")
        self.assertEqual(actual["start_session"], "2026-09-09")
        self.assertAlmostEqual(actual["portfolio_return"], 0)
        # Daily target rebalancing would instead produce 1.15 * .925 - 1.
        self.assertNotAlmostEqual(actual["portfolio_return"], 1.15*.925-1)

    def test_total_return_adjustment_does_not_add_dividends_twice(self):
        f = fixture()
        f["prices"]["series"]["AAA"]["bars"][1].update(dividend=5, split=2)
        result = validate(f)
        with_extra = evaluate(result, f["prices"], "2027-03-10T00:00:00Z")
        no_extra = evaluate(result, fixture()["prices"], "2027-03-10T00:00:00Z")
        self.assertEqual(with_extra["portfolio_return"], no_extra["portfolio_return"])
        f["prices"]["series"]["AAA"]["price_basis"] = "raw"
        with self.assertRaisesRegex(ValueError, "total_return_adjusted"):
            evaluate(result, f["prices"], "2027-03-10T00:00:00Z")

    def test_fx_conversion_and_fx_cost(self):
        f = fixture()
        f["decision"]["base_currency"] = "KRW"
        f["prices"]["fx"] = {"USD": dict(currency="USD",base_currency="KRW",provider="fixture",reference="fixture://fx",retrieved_at="2027-03-10T00:00:00Z",bars=[dict(date="2026-09-09",open=1000,close=1000),dict(date="2027-03-09",open=1200,close=1200)])}
        actual = evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")
        self.assertAlmostEqual(actual["contributions"]["AAA"], .15*(1.2*1.2/1.0015*.9985-1))
        self.assertEqual(actual["base_currency"], "KRW")

    def test_missing_terminal_and_suspended_assets_are_incomplete(self):
        for state in ("suspended", "delisted", "timeout", "credentials_missing"):
            f = fixture()
            f["prices"]["series"]["AAA"]["bars"][-1]["status"] = state
            actual = evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")
            self.assertEqual(actual["status"], "incomplete")
            self.assertNotIn("portfolio_return", actual)
        f = fixture()
        f["prices"]["series"]["AAA"]["bars"].pop()
        self.assertEqual(evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")["status"], "incomplete")

    def test_not_due_calendar_and_future_data(self):
        f = fixture()
        self.assertEqual(evaluate(validate(f), f["prices"], "2026-10-01T00:00:00Z")["status"], "not_due")
        f["decision"]["assets"][0]["calendar"] = "XKRX"
        self.assertEqual(evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")["issues"], ["mixed calendars unsupported"])
        f = fixture()
        f["prices"]["calendar"]["coverage_start"] = "2026-09-10T00:00:00Z"
        with self.assertRaisesRegex(ValueError, "calendar coverage"):
            evaluate(validate(f), f["prices"], "2027-03-10T00:00:00Z")
        f = fixture()
        with self.assertRaisesRegex(ValueError, "data as_of"):
            evaluate(validate(f), f["prices"], "2027-03-09T22:00:00Z")

    def test_append_only_revisions_and_thesis_independent_of_returns(self):
        f = fixture()
        result = validate(f)
        f["prices"]["thesis_reviews"] = {"AAA": dict(status="intact",reasoning="Operating thesis remains intact despite relative underperformance",as_of="2027-03-09T21:00:00Z",sources=["fixture://filing"])}
        with tempfile.TemporaryDirectory() as tmp:
            folder = save_snapshot(result, tmp)
            first = evaluate(result, f["prices"], "2027-03-10T00:00:00Z")
            path1 = save_evaluation(folder, first, "Initial evaluation")
            body = path1.read_bytes()
            f["prices"]["series"]["AAA"]["bars"][-1]["close"] = 121
            second = evaluate(result, f["prices"], "2027-03-10T00:00:00Z")
            path2 = save_evaluation(folder, second, "Provider corrected adjusted close")
            self.assertNotEqual(path1, path2)
            self.assertEqual(path1.read_bytes(), body)
            self.assertNotEqual(first["data_sha256"], second["data_sha256"])
            self.assertEqual(first["thesis_reviews"]["AAA"]["status"], "intact")
            self.assertLess(first["relative_return"], 0)
            self.assertEqual(load_snapshot(folder), result)


class IntegrationTests(unittest.TestCase):
    def test_risk_extra_etf_and_future_price_cutoff(self):
        import pandas as pd
        import build_risk_snapshot as risk_builder

        days = pd.bdate_range(end="2026-09-08", periods=280)
        rows = [{"date": str(day.date()), "close": 100 + i * .1} for i, day in enumerate(days)]
        future = {"date": "2026-09-09", "close": 999999}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root/"predict.json").write_text(json.dumps(fixture()["predict"]), encoding="utf-8")
            argv = ["build_risk_snapshot.py", "--predict-json", str(root/"predict.json"), "--extra-tickers", "MARKET", "--benchmark", "MARKET", "--output", str(root/"risk.json")]
            with patch.object(sys, "argv", argv), patch.object(risk_builder, "get_prices", return_value=rows+[future]) as prices, redirect_stdout(io.StringIO()):
                risk_builder.main()
            snapshot = json.loads((root/"risk.json").read_text(encoding="utf-8"))
            self.assertEqual(set(snapshot["annualized_volatility"]), {"AAA", "MARKET"})
            self.assertEqual(snapshot["market_regime"]["as_of_date"], "2026-09-08")
            self.assertLess(snapshot["annualized_volatility"]["MARKET"], .01)
            self.assertEqual({call.args[0] for call in prices.call_args_list}, {"AAA", "MARKET"})

    def test_cli_decision_to_report_snapshot_and_evaluation(self):
        f = fixture()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name in ("decision", "predict", "analyses", "risk", "prices"):
                (root/f"{name}.json").write_text(json.dumps(f[name]), encoding="utf-8")
            command = [sys.executable, str(SCRIPTS/"generate_portfolio_report.py"), "--predict-json", str(root/"predict.json"),
                       "--investor-json", str(root/"analyses.json"), "--risk-json", str(root/"risk.json"),
                       "--decision-json", str(root/"decision.json"), "--output-dir", str(root)]
            proc = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", env={**__import__('os').environ, "PYTHONIOENCODING": "utf-8"})
            self.assertEqual(proc.returncode, 0, proc.stderr)
            output = root/f"{f['decision']['decision_id']}.json"
            self.assertTrue(output.with_suffix('.xlsx').exists())
            result = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(result["weights"], {"AAA": .15, "MARKET": .6})
            snapshot = root/"decisions"/f["decision"]["decision_id"]
            command = [sys.executable, str(SCRIPTS/"evaluate_decision.py"), "--snapshot", str(snapshot), "--prices-json", str(root/"prices.json"),
                       "--as-of", "2027-03-10T00:00:00Z", "--reason", "Synthetic end-to-end test"]
            proc = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            evaluation = json.loads((snapshot/"evaluations/000001.json").read_text(encoding="utf-8"))
            self.assertEqual(evaluation["status"], "complete")


if __name__ == "__main__":
    unittest.main()
