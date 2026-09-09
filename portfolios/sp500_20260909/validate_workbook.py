"""Read the exported workbook and compare persisted cells with source JSON."""
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parents[1] / 'outputs/01a08599-dd56-71d0-8f8e-c239feafefd1/sp500_20260909_portfolio.xlsx'


def read(name):
    return json.loads((HERE / name).read_text(encoding='utf-8-sig'))


def main():
    portfolio = read('sp500_portfolio.json')
    predict = read('sp500_predict.json')
    investor = read('investor_analyses.json')
    values = load_workbook(WORKBOOK, data_only=True)
    formulas = load_workbook(WORKBOOK, data_only=False)
    assert values.sheetnames == ['포트폴리오', '전체 순위', '독립 투자자 분석', '위험과 근거']
    allocation = values.worksheets[0]
    actual = {allocation.cell(row, 1).value: allocation.cell(row, 4).value for row in range(9, 13)}
    assert actual == portfolio['weights']
    assert math.isclose(allocation['D13'].value, portfolio['cash_weight'])
    assert math.isclose(allocation['D14'].value, 1.0)
    assert formulas.worksheets[0]['D14'].value == '=SUM(D9:D13)'
    ranks = list(values.worksheets[1].iter_rows(min_row=7, max_row=509, min_col=1, max_col=8, values_only=True))
    assert len(predict['rankings']) == len(ranks) == 503
    for actual, expected in zip(ranks, predict['rankings']):
        assert actual[:4] == (expected['rank'], expected['ticker'], expected['company_name'], expected['total_score'])
    names = {'buffett': '버핏', 'lynch': '린치', 'fisher': '피셔'}
    signals = {'bullish': '긍정', 'neutral': '중립', 'bearish': '부정'}
    expected_rows = []
    for stock in predict['rankings'][:30]:
        for name in names:
            analysis = investor['analyses'][stock['ticker']][name]
            expected_rows.append((stock['ticker'], names[name], signals[analysis['signal']], analysis['confidence'], analysis['data_quality'], analysis['reasoning']))
    actual_rows = list(values.worksheets[2].iter_rows(min_row=7, max_row=96, min_col=1, max_col=6, values_only=True))
    assert len(actual_rows) == 90 and actual_rows == expected_rows
    error_cells = [f'{sheet.title}!{cell.coordinate}' for sheet in values for row in sheet for cell in row if cell.data_type == 'e']
    assert not error_cells, error_cells
    summary = {
        'checked_at_utc': datetime.now(timezone.utc).isoformat(),
        'workbook': str(WORKBOOK),
        'sha256': hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
        'status': 'pass',
        'sheet_count': len(values.sheetnames),
        'rankings_matched': len(ranks),
        'independent_analyses_matched': len(actual_rows),
        'weights_and_cached_total_matched': True,
        'formula_error_cells': error_cells,
        'native_excel_execution': 'not_performed',
    }
    (HERE / 'workbook_validation.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
