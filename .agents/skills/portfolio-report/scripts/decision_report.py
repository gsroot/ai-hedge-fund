"""Render every format from the same validated decision, including rejected proposals."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from portfolio_decision import canonical_bytes, save_snapshot, validate_decision


def report_rows(result):
    rows = [("판단", "상태", result["status"]), ("판단", "방식", result["methodology"]),
            ("판단", "성과 근거", result["performance_evidence"]["status"])]
    doc = result["decision"]
    rows += [("판단", key, doc[key]) for key in ("decision_id", "analysis_date", "created_at", "input_cutoff", "author", "base_currency", "benchmark")]
    if result["status"] in {"ready", "conditional"}:
        rows += [("최종 비중", ticker, weight) for ticker, weight in result["weights"].items()]
        rows.append(("최종 비중", "CASH", result["cash_weight"]))
    for key, value in doc["user_context"].items():
        rows.append(("사용자 맥락", key, value))
    for asset in doc["assets"]:
        rows.extend((asset["ticker"], key, asset[key]) for key in ("asset_type", "quote", "calendar"))
        rows.append((asset["ticker"], "선택/제외 이유", asset["selection_reason"]))
        rows.append((asset["ticker"], "검토한 비중", asset["weight"]))
        for key, value in asset.get("thesis", {}).items():
            rows.append((asset["ticker"], key, value))
        if "etf" in asset:
            rows.append((asset["ticker"], "ETF 자료", asset["etf"]))
    for category, values in (("현금 판단", doc["cash"]), ("제약 검사", result["constraint_checks"]),
                             ("적용 한도", doc["constraints"]),
                             ("보유 대비 판단", result["actions"]),
                             ("평가 규칙", doc["evaluation"]), ("독립 근거", result["evidence"]),
                             ("시장 참고값", result["market_regime"])):
        rows.extend((category, key, value) for key, value in values.items())
    rows += [("한계", str(i + 1), value) for i, value in enumerate(result["limitations"])]
    rows += [("출처", source["id"], source) for source in result["research_provenance"]["sources"]]
    return rows


def _rounded(value):
    if isinstance(value, float):
        return round(value, 10)
    if isinstance(value, dict):
        return {key: _rounded(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_rounded(item) for item in value]
    return value


def _cell(value):
    value = _rounded(value)
    return json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) or value is None else value


def is_percent(category, key):
    return (category == "최종 비중" or key == "검토한 비중"
            or (category == "현금 판단" and key in {"weight", "liquidity_weight", "annual_rate"})
            or (category == "시장 참고값" and key in {"target_cash_weight", "target_equity_weight", "cash_difference"})
            or (category == "제약 검사" and key == "lookthrough_coverage_weight")
            or (category == "적용 한도" and key.endswith("_weight")))


def display_rows(result):
    """Excel truncates long strings; split source/holdings details losslessly."""
    for category, key, value in report_rows(result):
        rendered = _cell(value)
        if isinstance(rendered, str) and len(rendered) > 30000:
            for start in range(0, len(rendered), 30000):
                yield category, f"{key} [{start // 30000 + 1}]", rendered[start:start + 30000]
        else:
            yield category, key, rendered


def render_text(result):
    lines = [f"현재 자료에 따른 추천: {result['decision_id']} ({result['analysis_date']})"]
    for category, key, value in display_rows(result):
        display = f"{value:.2%}" if is_percent(category, key) else str(_cell(value))
        lines.append(f"{category} | {key}: {display}")
    lines.append("현재 자료와 가정에 따른 연구 목적의 추천이며 미래 성과는 미확정입니다.")
    return "\n".join(lines) + "\n"


def write_reports(result, json_path, xlsx_path=None):
    json_path = Path(json_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_bytes(canonical_bytes(result))
    if json.loads(json_path.read_text(encoding="utf-8")) != result:
        raise ValueError("decision JSON round trip mismatch")
    text = render_text(result)
    json_path.with_suffix(".txt").write_text(text, encoding="utf-8")
    if xlsx_path is not None:
        workbook = Workbook()
        ws = workbook.active
        ws.title = "현재 추천"
        ws.append(["구분", "항목", "값"])
        for category, key, value in display_rows(result):
            # Strings from external sources must remain text, never Excel formulas.
            ws.append([category, key, _cell(value)])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if is_percent(category, key):
                ws.cell(ws.max_row, 3).number_format = "0.00%"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 32
        ws.column_dimensions["C"].width = 100
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(xlsx_path)
    return text


def run_decision_report(args, predict, analyses, risk, investors):
    decision = json.loads(Path(args.decision_json).read_text(encoding="utf-8"))
    result = validate_decision(decision, predict, analyses, risk, investors)
    output_dir = Path(args.output_dir)
    json_path = Path(args.portfolio_json) if args.portfolio_json else output_dir / f"{result['decision_id']}.json"
    snapshot_root = Path(args.snapshot_dir) if args.snapshot_dir else output_dir / "decisions"
    if json_path.resolve().is_relative_to(snapshot_root.resolve()):
        raise ValueError("presentation output must be outside immutable snapshot directory")
    # Reserve a new immutable identity before emitting mutable presentation copies.
    snapshot = save_snapshot(result, snapshot_root)
    text = write_reports(result, json_path, json_path.with_suffix(".xlsx") if args.xlsx.lower() in {"yes", "true", "1", "excel", "xlsx"} else None)
    print(text)
    print(f"추천 스냅샷: {snapshot}")
    return result
